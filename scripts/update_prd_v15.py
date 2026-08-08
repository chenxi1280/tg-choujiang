# ruff: noqa: E501, RUF001

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

SOURCE = Path("doc/TG群抽奖机器人PRD_V1.4_多奖品随机开奖版.docx")
TARGET = Path("doc/TG群抽奖机器人PRD_V1.5_极简创建版.docx")
CONTENT_WIDTH_DXA = "9360"
TABLE_INDENT_DXA = "120"

REPLACEMENTS = {
    "点击“创建抽奖”，选择机器人已加入且具备权限的目标群。": "点击“创建抽奖”，选择机器人已加入且具备权限的目标群；系统立即打开一张可反复编辑的活动配置面板。",
    "填写活动名称，设置预热开始、活动截止和开奖时间。": "必填项仅为活动标题、一次本地开奖时间和至少一个奖品。发布即开始预热，系统自动在开奖前 5 分钟截止报名并完成结算排空。",
    "打开与机器人的私聊，依次点击“创建抽奖 → 选择群组 → 设置奖品 → 设置时间 → 设置攒票规则 → 预览 → 发布到群”。发布前请重点核对：截止时间、开奖时间、邀请奖励、留存时长和频道门槛。": "打开与机器人的私聊，点击“创建抽奖 → 选择群组”进入配置面板。标题、开奖时间、奖品、领奖说明、活跃票、邀请票和关注频道均可独立编辑并返回面板，不要求按固定顺序完成。面板本身就是发布预览；缺少必填项时明确列出，不创建数据库半成品。",
    "管理员在私聊完成配置；服务端重新校验数据库授权、Telegram 管理员身份和所需权限。": "管理员在私聊通过单一配置面板逐项完成草稿；未发布数据保存于 24 小时 AdminSession。每个字段一次只接收一个简单输入，所有按钮使用用户与私聊绑定的一次性回调令牌。发布时服务端重新校验数据库授权、Telegram 管理员身份和所需权限。",
    "发布事务锁定 Campaign，冻结 config_revision，写入 prizes、规则快照、截止/开奖 scheduled_jobs 和群卡片 telegram_outbox；状态 draft -> publishing。": "点击发布时统一校验标题、开奖时间和奖品，按群时区把本地时间转换为带时区时间；以发布时刻生成 preheat_at，以开奖时间减 5 分钟生成 cutoff_at。随后同一发布事务锁定 Campaign，冻结 config_revision，写入 prizes、规则快照、生命周期 scheduled_jobs 和群卡片 telegram_outbox；状态 draft -> publishing。",
    "配置：管理员私聊输入“奖品名称 | 中奖人数 | 领取说明”，可连续添加多个奖品；系统按添加顺序生成 position。": "配置：管理员在奖品子面板先输入奖品名称，再单独输入正整数中奖人数；可连续添加、删除最后一个或清空。活动级领奖说明自动复制到每个 Prize，系统按添加顺序生成 position。",
}


def paragraphs(container: DocumentType | _Cell):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs(cell)


def replace_paragraph(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def apply_replacements(document: DocumentType) -> None:
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            replacement = REPLACEMENTS.get(paragraph.text)
            if replacement:
                replace_paragraph(paragraph, replacement)
            for run in paragraph.runs:
                run.text = run.text.replace("V1.4", "V1.5")


def restart_numbered_list(document: DocumentType, items: list[Paragraph]) -> None:
    numbering = document.part.numbering_part.element
    source_num_id = int(document.styles["List Number"].element.pPr.numPr.numId.val)
    source = next(
        num
        for num in numbering.findall(qn("w:num"))
        if int(num.get(qn("w:numId"))) == source_num_id
    )
    abstract_id = source.find(qn("w:abstractNumId")).get(qn("w:val"))
    next_id = max(int(num.get(qn("w:numId"))) for num in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(next_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    for item in items:
        num_pr = item._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = next_id


def normalize_table_geometry(document: DocumentType) -> None:
    for table in document.tables:
        properties = table._tbl.tblPr
        width = properties.first_child_found_in("w:tblW")
        if width is None:
            width = OxmlElement("w:tblW")
            properties.insert(0, width)
        width.set(qn("w:type"), "dxa")
        width.set(qn("w:w"), CONTENT_WIDTH_DXA)
        indent = properties.first_child_found_in("w:tblInd")
        if indent is None:
            indent = OxmlElement("w:tblInd")
            properties.insert(1, indent)
        indent.set(qn("w:type"), "dxa")
        indent.set(qn("w:w"), TABLE_INDENT_DXA)


def add_creation_section(document: DocumentType) -> None:
    document.add_heading("33. 极简面板式创建补充设计", level=1)
    document.add_heading("33.1 体验合同", level=2)
    document.add_paragraph(
        "创建体验采用“一个配置面板、字段独立编辑、随时返回”的结构。管理员不需要记忆步骤，也不需要使用竖线拼接多个字段；每次输入只回答当前一个问题。面板持续展示最终活动卡片所需的核心信息和发布缺口。"
    )
    flow_items = []
    for text in (
        "必填项只有活动标题、一次群组本地开奖时间和至少一个奖品。",
        "标题、开奖时间、奖品、领奖说明、活跃票、邀请票和关注频道均有独立入口；设置顺序不受限制。",
        "奖品采用名称与中奖人数两步录入，可连续添加多个不同奖品；领奖说明采用活动级默认值并可单独修改。",
        "活跃票和邀请票使用可读预设，不再要求输入六段技术参数；关注频道默认不限制。",
        "面板提供清空配置与取消创建。缺少必填项时点击发布只返回缺口，不创建 Campaign。",
    ):
        document.add_paragraph(text, style="List Bullet")
    document.add_heading("33.2 时间与开奖边界", level=2)
    document.add_paragraph(
        "V1.5 只简化创建方式，不改变开奖事实来源和算法。管理员输入的是对外公示的 draw_at；系统以发布事务时间生成 preheat_at，并以 draw_at - settlement_drain_seconds 生成 cutoff_at。默认排空窗口为 300 秒，因此开奖时间必须至少晚于当前时间 5 分钟。"
    )
    document.add_paragraph(
        "所有输入和展示均使用群组时区；数据库继续保存带时区时间。群活动卡片展示 YYYY-MM-DD HH:MM 的群组本地时间，不直接暴露 UTC 或 ISO 偏移量。"
    )
    document.add_paragraph(
        "截止后仍按原合同排空 Durable Inbox、复核频道资格与邀请留存、冻结票数快照和 prize_plan_hash，再使用 sha256-weighted-v2 按奖品顺位与名额执行加权随机抽取。"
    )
    document.add_heading("33.3 草稿与发布数据流", level=2)
    for text in (
        "选择群组：实时校验管理员身份，读取群名称和时区，创建 draft:idle AdminSession；草稿有效期 24 小时。",
        "编辑字段：按钮回调只改变会话 state；文本输入完成边界校验后合并进 AdminSession.data，并重新渲染同一面板消息。",
        "编辑奖品：draft:prize_name 保存临时名称，draft:prize_count 校验正整数后追加 prizes 数组；未完成的一半奖品不进入正式列表。",
        "发布校验：统一检查 title、draw_at、prizes，重新校验机器人邀请链接权限，并从面板数据构造 CampaignDraft、CampaignRules 和 PrizeSpec。",
        "持久化：只有发布校验通过后才创建 Campaign、RuleRevision 和 Prize；随后冻结规则及奖品方案，写 lifecycle Jobs 与 send_campaign_card Outbox。",
        "远端完成：群卡片 confirmed 并保存 remote_message_id 后，Campaign 才进入 scheduled/running；本地面板显示已提交不等于 Telegram 已送达。",
    ):
        flow_items.append(document.add_paragraph(text, style="List Number"))
    restart_numbered_list(document, flow_items)
    document.add_heading("33.4 状态、异常与验收", level=2)
    for text in (
        "会话状态仅包含 draft:idle、draft:title、draft:draw、draft:prize_name、draft:prize_count、draft:claim 和 draft:channel；任一子面板都可返回配置面板。",
        "非法标题、时间格式、过近开奖时间、非正整数中奖人数和非法频道 ID 必须保留当前输入状态并显示明确错误，允许原地重试。",
        "发布时权限被撤销、群不存在或同群已有非终态活动时必须失败关闭；不得清空草稿或伪造发布成功。",
        "验收必须覆盖乱序设置字段、多奖品不同人数、删除/清空、默认规则与预设切换、24 小时草稿、本地时区、必填缺口、发布权限复核和真实群卡片送达。",
        "生产验收继续分层报告：配置面板可用、数据库 Campaign/Prize 已提交、群卡片已确认、报名计票已发生、快照开奖已完成是五类不同事实。",
    ):
        document.add_paragraph(text, style="List Bullet")


def update_document() -> None:
    document = Document(SOURCE)
    apply_replacements(document)
    add_creation_section(document)
    normalize_table_geometry(document)
    document.save(TARGET)


if __name__ == "__main__":
    update_document()
