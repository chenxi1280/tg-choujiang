# ruff: noqa: E501, RUF001

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips
from update_prd_v15 import paragraphs, restart_numbered_list

SOURCE = Path("doc/TG群抽奖机器人PRD_V1.7_LotteryBestBot完整对标版.docx")
TARGET = Path("doc/TG群抽奖机器人PRD_V1.8_LotteryBestBot修复合同版.docx")
CONTRACT = Path("docs/development/v1.8-lotterybest-repair-prd.md")
SOURCE_SHA256 = "76897347261ac2a8979e4069425aefa482d21e6798388568c89ca0d2dd5f262e"

CJK_FONT = "Arial Unicode MS"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_VERTICAL_DXA = 80
CELL_MARGIN_HORIZONTAL_DXA = 120
HEADER_FILL = "D9E2F3"
METADATA_LABEL = "LotteryBestBot修复合同版"

SECTION_PATTERN = re.compile(r"^(\d+)\.\s+(.+)$")
SUBSECTION_PATTERN = re.compile(r"^(\d+)\.(\d+)\s+(.+)$")
NUMBERED_PATTERN = re.compile(r"^\d+\.\s+(.+)$")


def verify_source() -> None:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    if digest != SOURCE_SHA256:
        raise RuntimeError("V1.7 reference changed; re-distill the DOCX template")


def replace_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def apply_visible_version(document: DocumentType) -> None:
    for paragraph in document.paragraphs[:6]:
        if "V1.7 LotteryBestBot完整对标版" in paragraph.text:
            replace_paragraph_text(
                paragraph,
                paragraph.text.replace(
                    "V1.7 LotteryBestBot完整对标版",
                    "V1.8 LotteryBestBot修复合同版",
                ),
            )
    for section in document.sections:
        _replace_container_version(section.header)
        _replace_container_version(section.footer)
    _replace_version_scope(document)


def _replace_container_version(container) -> None:
    for paragraph in paragraphs(container):
        for run in paragraph.runs:
            run.text = run.text.replace("V1.7", "V1.8")
            run.text = run.text.replace("LotteryBestBot 完整对标版", METADATA_LABEL)
            run.text = run.text.replace("LotteryBestBot完整对标版", METADATA_LABEL)


def _replace_version_scope(document: DocumentType) -> None:
    if not document.tables:
        return
    for row in document.tables[0].rows:
        if row.cells[0].text.strip() != "版本范围":
            continue
        replace_paragraph_text(
            row.cells[1].paragraphs[0],
            "V1.8：V1.7 全量能力 + LotteryBestBot 修复合同 + Python/PostgreSQL 数据、事务、迁移与验收边界",
        )
        return


def plain_text(text: str) -> str:
    return text.replace("`", "").replace("**", "")


def renumber_heading(text: str, level: int) -> str:
    pattern = SECTION_PATTERN if level == 2 else SUBSECTION_PATTERN
    match = pattern.match(text)
    if match is None:
        return plain_text(text)
    if level == 2:
        return f"36.{int(match.group(1)) + 1} {plain_text(match.group(2))}"
    return f"36.{int(match.group(1)) + 1}.{match.group(2)} {plain_text(match.group(3))}"


def split_table_row(line: str) -> list[str]:
    return [plain_text(cell.strip()) for cell in line.strip().strip("|").split("|")]


def is_table_separator(cells: list[str]) -> bool:
    return all(cell.replace(":", "").replace("-", "") == "" for cell in cells)


def collect_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].startswith("|"):
        cells = split_table_row(lines[index])
        if not is_table_separator(cells):
            rows.append(cells)
        index += 1
    return rows, index


def table_widths(column_count: int) -> list[int]:
    if column_count == 2:
        return [5600, 3760]
    if column_count == 5:
        return [1100, 900, 3400, 2600, 1360]
    base = CONTENT_WIDTH_DXA // column_count
    widths = [base] * column_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def set_cell_margins(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        margin = (
            CELL_MARGIN_VERTICAL_DXA if edge in {"top", "bottom"} else CELL_MARGIN_HORIZONTAL_DXA
        )
        node.set(qn("w:w"), str(margin))
        node.set(qn("w:type"), "dxa")


def set_row_rules(row, header: bool) -> None:
    properties = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    properties.append(cannot_split)
    if header:
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        properties.append(repeat)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    _set_dimension(properties, "tblW", CONTENT_WIDTH_DXA)
    _set_dimension(properties, "tblInd", TABLE_INDENT_DXA)
    grid = table._tbl.tblGrid
    for node in list(grid):
        grid.remove(node)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    _set_table_cells(table, widths)


def _set_dimension(properties, tag: str, value: int) -> None:
    node = properties.first_child_found_in(f"w:{tag}")
    if node is None:
        node = OxmlElement(f"w:{tag}")
        properties.insert(0 if tag == "tblW" else 1, node)
    node.set(qn("w:type"), "dxa")
    node.set(qn("w:w"), str(value))


def _set_table_cells(table, widths: list[int]) -> None:
    for row_index, row in enumerate(table.rows):
        set_row_rules(row, row_index == 0)
        for column_index, cell in enumerate(row.cells):
            cell.width = Twips(widths[column_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell._tc.tcPr.tcW.set(qn("w:type"), "dxa")
            cell._tc.tcPr.tcW.set(qn("w:w"), str(widths[column_index]))
            set_cell_margins(cell)
            _format_table_cell(cell, row_index == 0, column_index)


def _format_table_cell(cell, header: bool, column_index: int) -> None:
    if header:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), HEADER_FILL)
        cell._tc.get_or_add_tcPr().append(shading)
    for paragraph in cell.paragraphs:
        if header or column_index in {0, 1, 4}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.bold = header


def add_table(document: DocumentType, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    set_table_geometry(table, table_widths(len(rows[0])))


def add_numbered_block(document: DocumentType, lines: list[str], start: int) -> int:
    items = []
    index = start
    while index < len(lines):
        match = NUMBERED_PATTERN.match(lines[index])
        if match is None:
            break
        items.append(document.add_paragraph(plain_text(match.group(1)), style="List Number"))
        index += 1
    restart_numbered_list(document, items)
    return index


def add_contract(document: DocumentType) -> None:
    lines = CONTRACT.read_text(encoding="utf-8").splitlines()
    start = next(index for index, line in enumerate(lines) if line.startswith("## 0."))
    document.add_page_break()
    document.add_heading("36. V1.8 LotteryBestBot 修复合同", level=1)
    index = start
    while index < len(lines):
        index = add_contract_line(document, lines, index)


def add_contract_line(document: DocumentType, lines: list[str], index: int) -> int:
    line = lines[index]
    if not line:
        return index + 1
    if line.startswith("### "):
        document.add_heading(renumber_heading(line[4:], 3), level=3)
        return index + 1
    if line.startswith("## "):
        document.add_heading(renumber_heading(line[3:], 2), level=2)
        return index + 1
    if line.startswith("| "):
        rows, next_index = collect_table(lines, index)
        add_table(document, rows)
        return next_index
    if NUMBERED_PATTERN.match(line):
        return add_numbered_block(document, lines, index)
    if line.startswith("- "):
        document.add_paragraph(plain_text(line[2:]), style="List Bullet")
        return index + 1
    document.add_paragraph(plain_text(line))
    return index + 1


def normalize_fonts(document: DocumentType) -> None:
    for style in document.styles:
        if style.type != 1:
            continue
        style.font.name = CJK_FONT
        _set_font_names(style.element.get_or_add_rPr().get_or_add_rFonts())
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            for run in paragraph.runs:
                run.font.name = CJK_FONT
                _set_font_names(run._element.get_or_add_rPr().get_or_add_rFonts())


def mark_data_table_headers(document: DocumentType) -> None:
    for table in document.tables:
        if len(table.rows) < 2 or len(table.columns) < 2:
            continue
        properties = table.rows[0]._tr.get_or_add_trPr()
        if properties.find(qn("w:tblHeader")) is not None:
            continue
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        properties.append(repeat)


def _set_font_names(fonts) -> None:
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), CJK_FONT)


def update_document() -> None:
    verify_source()
    document = Document(SOURCE)
    apply_visible_version(document)
    add_contract(document)
    normalize_fonts(document)
    mark_data_table_headers(document)
    document.core_properties.title = "TG 群抽奖机器人 V1.8 LotteryBestBot 修复合同版"
    document.save(TARGET)


if __name__ == "__main__":
    update_document()
