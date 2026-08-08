# ruff: noqa: E501, RUF001

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from update_prd_v15 import (
    normalize_table_geometry,
    paragraphs,
    replace_paragraph,
    restart_numbered_list,
)

SOURCE = Path("doc/TG群抽奖机器人PRD_V1.5_极简创建版.docx")
TARGET = Path("doc/TG群抽奖机器人PRD_V1.6_极简交互重构版.docx")
CJK_FONT = "Noto Sans SC"

REPLACEMENTS = {
    "创建体验采用“一个配置面板、字段独立编辑、随时返回”的结构。管理员不需要记忆步骤，也不需要使用竖线拼接多个字段；每次输入只回答当前一个问题。面板持续展示最终活动卡片所需的核心信息和发布缺口。": "创建体验采用“/new 直接进入、首次引导、最终卡片预览”的结构。管理员不需要先经过首页和群组选择层，也不需要理解活跃票、邀请票等技术配置；首次只依次回答奖品名称、中奖人数、必要时的活动群、标题和开奖时间。完成后进入最终卡片预览，所有字段可独立修改。",
    "必填项只有活动标题、一次群组本地开奖时间和至少一个奖品。": "必填项只有活动群、抽奖标题、一次群组本地开奖时间和至少一个奖品；若管理员只管理一个群，系统自动绑定该群。",
    "标题、开奖时间、奖品、领奖说明、活跃票、邀请票和关注频道均有独立入口；设置顺序不受限制。": "最终预览页只提供活动群、抽奖标题、奖品内容、开奖时间、领奖说明和中奖规则六个业务入口，并以两行三列紧凑展示；活跃与邀请规则合并为中奖规则，不再占用首页按钮。",
    "面板提供清空配置与取消创建。缺少必填项时点击发布只返回缺口，不创建 Campaign。": "预览页提供发布抽奖、清空配置和返回首页；清空必须二次确认。缺少必填项时预览直接列出缺口，点击发布只返回缺口，不创建 Campaign。",
}


def apply_updates(document: DocumentType) -> None:
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            replacement = REPLACEMENTS.get(paragraph.text)
            if replacement:
                replace_paragraph(paragraph, replacement)
            for run in paragraph.runs:
                run.text = run.text.replace("V1.5", "V1.6")
                run.text = run.text.replace("数据流转设计版", "极简交互重构版")


def add_v16_section(document: DocumentType) -> None:
    document.add_heading("34. 极简创建二次重构", level=1)
    document.add_heading("34.1 参考原则与差异边界", level=2)
    document.add_paragraph(
        "V1.6 学习极简抽奖机器人的创建信息架构：短入口、一步一问、最终卡片预览、三列业务按钮和明确的返回层级。仅学习交互，不复制其开奖合同。本项目继续以群活跃票与有效邀请票形成权重，使用不可变快照和 sha256-weighted-v2 完成可复算随机开奖。"
    )
    for text in (
        "私聊发送 /new 直接创建或继续未完成草稿；/start 仍展示首页。",
        "仅有一个可管理群时自动绑定；多个群时只在确实需要时展示群选择页。",
        "首次引导顺序为奖品名称 → 中奖人数 → 活动群（如需）→ 抽奖标题 → 快捷开奖时间 → 最终预览。",
        "设置子页只保留当前动作与返回上一层，禁止把技术参数重新铺回主面板。",
        "快捷开奖时间提供 1 小时后、今晚 20:00、明晚 20:00 和自定义时间。",
    ):
        document.add_paragraph(text, style="List Bullet")
    document.add_heading("34.2 页面与状态合同", level=2)
    items = []
    for text in (
        "入口：/new 查询 AdminSession。存在 draft:* 草稿则恢复并显示最终预览；不存在则创建会话并进入 prize_name。",
        "奖品：prize_name 只接收名称，prize_count 只接收正整数人数；完成后原子追加到 prizes，临时名称不进入正式奖品数组。",
        "群组：查询当前管理员可管理群；单群自动绑定，多群通过一次性安全回调绑定 chat_id、group_title 与 timezone。",
        "标题与时间：标题长度 1–80；快捷或自定义时间统一转换为带群时区的 draw_at，且必须晚于当前时间至少 5 分钟。",
        "预览：以最终抽奖卡片语义展示六个业务字段和缺口；按钮为两行三列，随后是发布、清空/首页。",
        "发布：重新校验管理员和邀请链接权限，构造 CampaignDraft 与 PrizeSpec；成功后清理 AdminSession，失败则保留草稿。",
    ):
        items.append(document.add_paragraph(text, style="List Number"))
    restart_numbered_list(document, items)

    document.add_heading("34.3 数据流转", level=2)
    document.add_paragraph(
        "Telegram Update → 私聊命令路由 → AdminSession(state, data, panel_message_id) → 字段校验与不可变合并 → 同一消息重新渲染 → 发布期权限复核 → Campaign/RuleRevision/Prize 事务 → lifecycle ScheduledJob 与 send_campaign_card Outbox → Telegram confirmed remote_message_id → Campaign scheduled/running。"
    )
    document.add_paragraph(
        "创建期数据只存在 AdminSession，不提前写 Campaign。正式发布后，计票仍经 Update Inbox、Activity/Invitation 规则、TicketLedger 唯一键进入账本；截止后冻结 DrawSnapshot 与 prize_plan_hash，再由原有加权随机算法生成 Winner。界面简化不改变任何票据、快照、幂等或开奖审计字段。"
    )

    document.add_heading("34.4 验收标准", level=2)
    for text in (
        "发送 /new 后不得先显示首页；新草稿直接询问奖品，旧草稿直接恢复最终预览。",
        "单群管理员不出现群选择步骤；多群管理员只有在缺少活动群时进入选择页。",
        "最终预览必须出现六个业务字段、缺口状态、两行三列编辑按钮、发布、二次确认清空和返回首页。",
        "可为每次抽奖设置不同奖品和不同中奖人数，并可继续添加、删除最后一个或清空奖品。",
        "创建交互变化不得改变群活跃计票、邀请留存、不可变快照、奖品顺位、多名额去重和加权随机算法。",
        "自动化必须覆盖 /new 直达、预览文案、三列按钮、缺口校验、多奖品人数和原有开奖回归；线上验收需单独验证 Telegram 实际回复。",
    ):
        document.add_paragraph(text, style="List Bullet")


def normalize_fonts(document: DocumentType) -> None:
    for style in document.styles:
        if style.type != 1:
            continue
        style.font.name = CJK_FONT
        _set_font_names(style.element.get_or_add_rPr().get_or_add_rFonts())
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            for run in paragraph.runs:
                run.font.name = CJK_FONT
                _set_font_names(run._element.get_or_add_rPr().get_or_add_rFonts())


def _set_font_names(fonts) -> None:
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), CJK_FONT)


def update_document() -> None:
    document = Document(SOURCE)
    apply_updates(document)
    add_v16_section(document)
    normalize_fonts(document)
    normalize_table_geometry(document)
    document.save(TARGET)


if __name__ == "__main__":
    update_document()
