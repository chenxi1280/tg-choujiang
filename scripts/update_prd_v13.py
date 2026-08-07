# ruff: noqa: E501, RUF001

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.table import _Cell
from docx.text.paragraph import Paragraph

SOURCE = Path("doc/TG群抽奖机器人PRD_V1.2_数据流转设计版.docx")
TARGET = Path("doc/TG群抽奖机器人PRD_V1.3_实现对齐版.docx")

PARAGRAPH_REPLACEMENTS = {
    "每次 callback_query 都必须调用 answerCallbackQuery；后台处理较长时先确认按钮，再展示最终结果。": (
        "每次 callback_query 到达后直接尽快调用 answerCallbackQuery；该交互 ACK 不进入业务 "
        "Outbox，也不代表业务成功。发布、邀请链接与开奖通知等可审计副作用仍必须走 Outbox。"
    ),
    "生产至少具备 PostgreSQL 备份、Alembic 记录、app/worker heartbeat、失败 Job 与 UNKNOWN delivery 告警。": (
        "生产每次迁移前生成 PostgreSQL custom-format 备份并执行保留策略；同时保留 Alembic "
        "记录、app/worker heartbeat、失败 Job、失败 Update 与 UNKNOWN delivery 告警。"
    ),
    "回调按钮的 answerCallbackQuery 通过 outbox 发送；交互提示不替代数据库业务事务和后续远端完成事实。": (
        "回调按钮先直接发送 answerCallbackQuery 交互 ACK，避免排队超时；ACK 不替代数据库业务事务。"
        "群卡片、邀请链接、开奖结果和中奖私聊仍通过 Outbox 留存远端完成事实。"
    ),
    "worker 只用 FOR UPDATE SKIP LOCKED 领取 telegram_outbox/scheduled_jobs，并以数据库时钟、lease_owner、lease_until 和 attempt_no 控制并发。": (
        "worker 只用 FOR UPDATE SKIP LOCKED 领取 telegram_outbox/scheduled_jobs，并以数据库时钟、"
        "lease_owner、lease_until 和 attempt_no 控制并发；迟到 Worker 必须再次校验 owner/state，"
        "不得覆盖租约接管后的状态。"
    ),
    "调用前先提交 DeliveryAttempt.gateway_started_at 与 request_hash；调用后另起事务追加 confirmed/failed/unknown 结果。": (
        "调用前先在 Outbox 提交 gateway_started_at 与 request_hash；调用后另起事务追加 "
        "DeliveryAttempt=confirmed/failed/unknown，再投影业务状态。"
    ),
}


def paragraphs(container: DocumentType | _Cell):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs(cell)


def replace_paragraph(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def update_document() -> None:
    document = Document(SOURCE)
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            replacement = PARAGRAPH_REPLACEMENTS.get(paragraph.text)
            if replacement:
                replace_paragraph(paragraph, replacement)
            for run in paragraph.runs:
                run.text = run.text.replace("V1.2", "V1.3")
    document.save(TARGET)


if __name__ == "__main__":
    update_document()
