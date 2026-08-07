# ruff: noqa: E501, RUF001

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

SOURCE = Path("doc/TG群抽奖机器人PRD_V1.3_实现对齐版.docx")
TARGET = Path("doc/TG群抽奖机器人PRD_V1.4_多奖品随机开奖版.docx")
DOCUMENT_FONT = "Arial Unicode MS"

REPLACEMENTS = {
    "配置奖品名称、数量、领取说明。": "逐个配置奖品名称、中奖人数、领取说明和抽取顺序；每期活动使用独立奖品方案。",
    "管理员配置活动，系统保存活动规则快照与奖品信息。": "管理员配置活动，系统保存本期独立的活动规则和奖品方案；发布时冻结 prize_plan_hash。",
    "活动截止时转入结算中，排空截止前事件并完成资格复核；条件通过后冻结参与者与票数，形成开奖快照。": "活动截止时转入结算中，排空截止前事件并完成资格复核；条件通过后将参与者、票数、规则版本和 prize_plan_hash 一并写入不可变开奖快照。",
    "开奖服务按快照中的票数加权抽取并提交中奖记录；群内公示和私聊通知由事务后的 outbox 独立送达。": "开奖服务按奖品顺位和各奖品中奖人数逐槽位加权随机抽取并提交中奖记录；默认中奖者从后续槽位候选池移除。群内公示、未分配名额和含领取说明的中奖私聊由事务后的 outbox 独立送达。",
    "开奖按票数加权、按奖品顺序抽取；默认同一人一期最多中一次。候选人数少于奖品份数时，剩余份数标记未派出。": "开奖按冻结票数加权，并按奖品顺位、奖品内名额编号依次抽取；默认同一人一期最多中一次。奖品顺位也是候选不足时的分配优先级；候选人数少于总中奖人数时，后续槽位标记未分配并在群内公告。",
    "记录 algorithm_version、快照哈希、随机输入与每次抽取区间。V1 提供可复核审计，但不宣称第三方密码学随机信标。": "记录 algorithm_version、候选快照哈希、prize_plan_hash、随机输入、奖品槽位与每次抽取区间。sha256-weighted-v2 使用确定性拒绝采样消除取模偏差；提供可复核审计，但不宣称第三方密码学随机信标。",
    "自动或手动开奖都以稳定 draw business_key 进入同一 DrawService；锁定 ready_to_draw Campaign 并校验 snapshot_hash。": "自动或手动开奖都以稳定 draw business_key 进入同一 DrawService；锁定 ready_to_draw Campaign，并校验 snapshot_hash、Campaign.prize_plan_hash、DrawSnapshot.prize_plan_hash 与当前奖品方案四者一致。",
    "V1 使用冻结累计权重加权抽取；中奖者立即从后续轮次候选池移除，保证同一期最多中一次。": "V1.4 使用冻结累计权重与 sha256-weighted-v2 拒绝采样逐槽位抽取；中奖者立即从后续槽位候选池移除，默认保证同一期最多中一次。",
    "开奖后核对 Snapshot total_weight、DrawRound 命中区间、Winner 数量和奖品槽位；数据库通过与 Telegram 送达分别报告。": "开奖后核对 Snapshot total_weight/prize_plan_hash、DrawRound 的 round_number/prize_slot_no/命中区间、Winner 数量及未分配槽位；数据库通过与 Telegram 群公告、中奖私聊送达分别报告。",
    "名称、数量、领取说明、排序": "名称、中奖人数、领取说明、抽取顺位",
    "名称、数量、领取说明": "名称、中奖人数、领取说明",
}


def paragraphs(container: DocumentType | _Cell):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs(cell)


def replace_paragraph(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def apply_replacements(document: DocumentType) -> None:
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            replacement = REPLACEMENTS.get(paragraph.text)
            if replacement:
                replace_paragraph(paragraph, replacement)
            for run in paragraph.runs:
                run.text = run.text.replace("V1.3", "V1.4")


def apply_portable_font(document: DocumentType) -> None:
    for style in document.styles:
        if style.type != 1:
            continue
        style.font.name = DOCUMENT_FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), DOCUMENT_FONT)
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            for run in paragraph.runs:
                run.font.name = DOCUMENT_FONT
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), DOCUMENT_FONT)


def add_design_section(document: DocumentType) -> None:
    document.add_heading("32. 每期多奖品与随机开奖补充设计", level=1)
    document.add_heading("32.1 配置合同", level=2)
    document.add_paragraph(
        "每期活动独立配置一个或多个奖品。每个奖品必须包含名称、中奖人数、领取说明和唯一抽取顺位；中奖人数必须大于 0。同一期不得出现重复顺位。发布预览和群活动卡片必须展示每个奖品的中奖人数、顺位及总中奖人数。"
    )
    document.add_paragraph(
        "发布成功后奖品方案不可原地修改。系统对奖品 ID、顺位、名称、中奖人数和领取说明生成 prize_plan_hash，并在 Campaign 与 DrawSnapshot 中双重保存；任一字段漂移都阻断结算或开奖。"
    )
    document.add_heading("32.2 数据流转", level=2)
    for text in (
        "配置：管理员私聊输入“奖品名称 | 中奖人数 | 领取说明”，可连续添加多个奖品；系统按添加顺序生成 position。",
        "发布：同一事务冻结规则版本和 prize_plan_hash，写生命周期 Job 与活动卡片 Outbox；Telegram 确认卡片送达后活动才进入运行状态。",
        "结算：排空截止前本活动 Update、邀请验证和资格复核后，按票数账本聚合候选权重，并把 prize_plan_hash 写入 DrawSnapshot 和 snapshot_hash 输入。",
        "开奖：按 position 从小到大、每个奖品按 prize_slot_no 从 1 到中奖人数执行；每槽位使用 random_input、round_number 和 nonce 的 SHA-256 拒绝采样选取权重区间。",
        "去重：默认 unique_winner=true，中奖者立即从后续全部奖品槽位移除；因此奖品顺位决定候选不足时的优先级。",
        "结果：每个计划槽位都生成 DrawRound；有候选人时生成 Winner，无候选人时 undistributed=true。群公告必须展示中奖名单和未分配名额，中奖私聊必须带本奖品领取说明。",
        "审计：DrawRun、DrawRound、Winner、群公告 Outbox 和中奖私聊 Outbox 同事务提交；远端通知失败不改变数据库中奖事实。",
    ):
        document.add_paragraph(text, style="List Number")
    document.add_heading("32.3 数据模型增量", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(
        table.rows[0].cells, ("对象", "新增/调整字段", "约束与用途"), strict=True
    ):
        cell.text = text
    rows = (
        ("Campaign", "prize_plan_hash", "发布时冻结本期奖品方案"),
        ("Prize", "winner_count", "> 0；campaign_id + position 唯一"),
        ("DrawSnapshot", "prize_plan_hash", "与候选快照共同构成开奖输入"),
        ("DrawRound", "prize_slot_no", "> 0；draw_run + prize + slot 唯一"),
        ("Outbox", "undistributed / claim_instructions", "披露空缺名额并告知领奖方式"),
    )
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    document.add_heading("32.4 异常与验收口径", level=2)
    for text in (
        "总中奖人数大于合格候选人数：不允许重复中奖补足，后续槽位保持未分配并公开数量。",
        "奖品方案发布后发生漂移：结算或开奖失败并暴露明确错误，不生成部分结果。",
        "重复自动/手动开奖：返回同一 algorithm_version 对应的 DrawRun，不产生第二组随机输入。",
        "验收必须覆盖至少两个奖品、不同中奖人数、候选不足、中奖者全局去重、奖品槽位唯一、领取说明私聊和奖品哈希漂移阻断。",
        "生产验收仍需真实 Telegram 群完成创建、报名、发票、快照、开奖、群公告和中奖私聊；迁移、健康检查与本地测试不等于业务闭环。",
    ):
        document.add_paragraph(text, style="List Bullet")


def update_document() -> None:
    document = Document(SOURCE)
    apply_replacements(document)
    add_design_section(document)
    apply_portable_font(document)
    document.save(TARGET)


if __name__ == "__main__":
    update_document()
