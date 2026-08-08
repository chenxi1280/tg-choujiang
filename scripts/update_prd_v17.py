# ruff: noqa: E501, RUF001

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from update_prd_v15 import normalize_table_geometry, paragraphs

SOURCE = Path("doc/TG群抽奖机器人PRD_V1.6_极简交互重构版.docx")
TARGET = Path("doc/TG群抽奖机器人PRD_V1.7_LotteryBestBot完整对标版.docx")
CJK_FONT = "Arial Unicode MS"


def apply_version(document: DocumentType) -> None:
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            for run in paragraph.runs:
                run.text = run.text.replace("V1.6", "V1.7")
                run.text = run.text.replace("极简交互重构版", "LotteryBestBot完整对标版")


def add_decision(document: DocumentType) -> None:
    document.add_heading("35. LotteryBestBot 完整功能对标", level=1)
    document.add_heading("35.1 产品决策与边界", level=2)
    document.add_paragraph(
        "V1.7 对标 @LotteryBestBot 已实测确认的命令、抽奖管理、个人地区、参加记录、草稿恢复和创建配置体验。对标的是用户功能与信息架构，不是随机抽取内核。本项目继续使用群活跃票、有效邀请票、不可变账本、资格复核、DrawSnapshot、多奖品顺位和 sha256-weighted-v2。"
    )
    document.add_paragraph(
        "任何参考页面若未完成只读实机观察，只能作为待验证项；不得据此宣称参考行为已确认。V1.7 的实现合同以本节和仓库 docs/development/v1.7-lotterybest-parity.md 为准。"
    )


def add_commands(document: DocumentType) -> None:
    document.add_heading("35.2 一级命令与管理首页", level=2)
    for text in (
        "/start 显示使用指南和主菜单，公开 /index、/new、/join、/region、/myid。",
        "/index 展示累计创建、进行中、已开奖、已取消数量，并提供活动分类、发起抽奖、添加群组和添加频道入口。",
        "/new 新草稿直接进入奖品名称；已有 draft:* 草稿时直接恢复最终预览。",
        "/join 只读展示当前用户参加过的抽奖，并显示标题、状态、报名资格、总票数和开奖时间；空状态明确返回未参与任何抽奖。",
        "/region 设置省级归属地区；首次设置后 5 分钟内可修改，随后锁定 30 天；每次变更写旧值、新值和时间审计。",
        "/myid 只返回当前 Telegram 用户 ID。",
    ):
        document.add_paragraph(text, style="List Bullet")


def add_creator_contract(document: DocumentType) -> None:
    document.add_heading("35.3 创建预览与字段合同", level=2)
    document.add_paragraph(
        "最终预览展示抽奖标题、封面、用户名限制、地区限制、参与条件、奖品、抽奖说明、人品加成、开奖方式和发奖人员。未设置必填项统一显示【等待设置】，默认用户名可为空、地区全国、人品加成关闭、发奖人员为创建者。"
    )
    for text in (
        "按钮布局：活动标题/上传封面/用户限制；参与条件/奖品内容/抽奖说明；开奖方式/发奖人员/人品加成；消息提示/发布抽奖；清空配置/载入上次抽奖；返回首页。",
        "封面只接收 Telegram 图片并保存 file_id；抽奖说明与领奖说明分离；发奖人员默认创建者，可输入用户 ID 或用户名。",
        "用户限制支持用户名必填开关和最多 3 个省级地区；报名时依据用户 /region 资料校验。",
        "参与条件必须包含一个已登记主群；额外条件逐个输入公开 @username 或数字 ID，发布时解析并冻结 chat ID、标题、类型、用户名，且 Bot 必须是该群或频道管理员。",
        "奖品按名称→正整数名额采集，支持多奖品，添加顺序就是开奖顺位。",
        "每位合格报名者固定获得 1 张幂等基础票；人品加成关闭时不发奖励票，标准和高活跃模式才按各自预设发放活跃票与邀请票。",
        "人品加成不得提供指定中奖或暗箱能力；页面明确每张票等概率、总票数越多权重越高。",
        "消息提示只控制扩展说明；关键失败、资格不通过、参与结果和开奖结果始终明确发送。",
        "发布成功保存上次配置 24 小时；载入上次只生成新草稿，不复用旧 Campaign、消息 ID 或生命周期任务。",
    ):
        document.add_paragraph(text, style="List Bullet")


def add_draw_contract(document: DocumentType) -> None:
    document.add_heading("35.4 开奖方式与本项目算法", level=2)
    document.add_paragraph(
        "定时开奖只要求一次群时区本地时间。满人开奖只要求期望参与人数，并冻结 30 天兜底开奖时刻。达到门槛后立即进入本项目的安全截止、Inbox 排空、资格复核、快照和开奖流程；这里的“立即”不允许绕过结算安全门。"
    )
    document.add_paragraph(
        "两种调度入口最终都生成唯一不可变 DrawSnapshot，保存配置 revision、账本水位、权重和 snapshot hash，并由 sha256-weighted-v2 按奖品顺位分配 Winner。参考产品的随机实现不进入本项目。"
    )


def add_data_and_acceptance(document: DocumentType) -> None:
    document.add_heading("35.5 数据流转与验收", level=2)
    document.add_paragraph(
        "Telegram Update → 私聊命令/安全回调 → UserProfile 或 AdminSession → 字段边界校验 → 最终预览 → 发布期管理员与可验证性复核 → Campaign.presentation + RuleRevision + Prize → ScheduledJob/Outbox → Telegram confirmed remote fact。报名后再经资格、TicketLedger、Settlement、DrawSnapshot、DrawRun 和 Winner。"
    )
    document.add_paragraph(
        "UserProfile 保存地区选择时间、5 分钟可编辑截止、30 天锁定截止和最近配置；UserRegionChange 保存每次地区变更审计。Campaign.presentation 冻结封面、说明、发奖人员、消息提示和额外参与条件。TicketLedger 的 registration 来源保存每位合格报名者唯一基础票。"
    )
    for text in (
        "自动化覆盖全部一级命令、地区 5 分钟/30 天、管理统计、参加列表、草稿恢复、载入上次和每个创建字段。",
        "迁移必须通过 upgrade/downgrade/upgrade 和模型漂移检查；旧规则缺失新增字段时使用显式 schema 默认值。",
        "发布必须走 master→release→GitHub Actions，并回读 deployed SHA、迁移 head、容器健康和 RestartCount。",
        "Telegram E2E 逐项点击本项目 Bot 的创建按钮，但不得为了验收发布真实测试抽奖。",
        "没有受控群时，报名、满人触发、计票、邀请和开奖继续标记 unproven，不由本地测试或健康检查替代。",
    ):
        document.add_paragraph(text, style="List Bullet")


def normalize_fonts(document: DocumentType) -> None:
    for style in document.styles:
        if style.type != 1:
            continue
        style.font.name = CJK_FONT
        _set_font_names(style.element.get_or_add_rPr().get_or_add_rFonts())
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            for run in paragraph.runs:
                run.font.name = CJK_FONT
                _set_font_names(run._element.get_or_add_rPr().get_or_add_rFonts())


def _set_font_names(fonts) -> None:
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), CJK_FONT)


def update_document() -> None:
    document = Document(SOURCE)
    apply_version(document)
    add_decision(document)
    add_commands(document)
    add_creator_contract(document)
    add_draw_contract(document)
    add_data_and_acceptance(document)
    normalize_fonts(document)
    normalize_table_geometry(document)
    document.save(TARGET)


if __name__ == "__main__":
    update_document()
